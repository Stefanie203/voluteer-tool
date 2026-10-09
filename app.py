import streamlit as st
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from io import BytesIO
import re

st.set_page_config(page_title="志愿学时自动统计工具", layout="wide")
st.title("📊 志愿学时自动统计工具")
st.caption("上传报名表 → 自动统计学时 → 生成符合规范的 Excel")

# ==================== 侧边栏配置（简化） ====================
with st.sidebar:
    st.header("⚙️ 配置")
    
    # 新增：周数输入框
    week_num = st.text_input("当前是第几周？（填入数字，如 2）", value="1")
    
    st.info("单次活动学时固定为 2 学时。系统将自动按 2、4、6、8... 的偶数进行分组。")
    
    st.subheader("地点归类映射")
    st.caption("已内置自动规则：地名以“号站”结尾，自动归为“宿管站”。")
    default_mapping = "图书馆=图书馆"
    mapping_text = st.text_area("额外映射（可选，每行一个）", value=default_mapping, height=100)

    location_map = {}
    for line in mapping_text.strip().split("\n"):
        if "=" in line:
            k, v = line.split("=", 1)
            location_map[k.strip()] = v.strip()
    
    # 新增：默认地点，防止出现“未知地点”
    default_location = st.text_input("如果表格没有地点列，默认按此处理", value="图书馆")

# ==================== 数据清洗函数 ====================
def clean_name(name):
    """去除名字前的班级前缀，例如 大数据2401卜晨 -> 卜晨"""
    if pd.isna(name): return None
    cleaned = re.sub(r'^[\u4e00-\u9fa5]{2,4}\d{4}', '', str(name)).strip()
    return cleaned if cleaned else str(name).strip()

# ==================== 数据输入 ====================
uploaded_file = st.file_uploader("上传 CSV / Excel 报名表", type=["csv", "xlsx", "xls"])

if uploaded_file is not None:
    try:
        raw_df = pd.read_csv(uploaded_file) if uploaded_file.name.endswith(".csv") else pd.read_excel(uploaded_file)
        st.success(f"✅ 文件读取成功，共 {len(raw_df)} 行")
    except Exception as e:
        st.error(f"文件读取失败：{e}")
        st.stop()

    st.subheader("🔍 原始数据预览")
    st.dataframe(raw_df.head(10), use_container_width=True)

    st.subheader("🛠️ 数据智能提取")
    
    name_cols = [c for c in raw_df.columns if "姓名" in str(c) or "志愿者" in str(c)]
    id_cols = [c for c in raw_df.columns if "学号" in str(c) or str(c).startswith("Unnamed")]
    loc_col = next((c for c in raw_df.columns if "地点" in str(c) or "站" in str(c) or "图书馆" in str(c)), None)

    if not name_cols or not id_cols:
        st.warning("未能完美自动识别，请手动选择对应的列：")
        cols = list(raw_df.columns)
        c1, c2 = st.columns(2)
        name_col = c1.selectbox("姓名列", cols, index=0)
        id_col = c2.selectbox("学号列", cols, index=min(1, len(cols) - 1))
        name_cols = [name_col]
        id_cols = [id_col]
    
    all_records = []
    for nc in name_cols:
        prefix = str(nc).replace("姓名", "").replace("志愿者", "")
        ic = prefix + "学号"
        if ic not in raw_df.columns:
            idx = raw_df.columns.get_loc(nc)
            if idx + 1 < len(raw_df.columns):
                ic = raw_df.columns[idx + 1]
        
        if ic in raw_df.columns:
            tmp = raw_df[[nc, ic]].copy()
            tmp.columns = ["姓名", "学号"]
            if loc_col and loc_col in raw_df.columns:
                tmp["地点"] = raw_df[loc_col].ffill()
            else:
                tmp["地点"] = default_location if default_location else "未知地点"
            tmp = tmp.dropna(subset=["姓名", "学号"])
            all_records.append(tmp)
    
    if not all_records:
        st.error("未能提取到有效数据，请检查表格格式。")
        st.stop()
        
    long_df = pd.concat(all_records, ignore_index=True)

    # 数据清洗
    long_df["姓名"] = long_df["姓名"].apply(clean_name)
    long_df["学号"] = long_df["学号"].astype(str).str.strip().str.replace(r"\.0$", "", regex=True)
    long_df["地点"] = long_df["地点"].astype(str).str.strip()
    
    long_df = long_df[(long_df["姓名"] != "") & (long_df["学号"] != "")]

    st.subheader("✏️ 数据修正（可直接双击编辑，已去除班级前缀）")
    edited = st.data_editor(long_df, num_rows="dynamic", use_container_width=True)

    # 自动映射地点
    def map_location(loc):
        if "号站" in str(loc):
            return "宿管站"
        return location_map.get(str(loc), str(loc))
    
    edited["活动分类"] = edited["地点"].apply(map_location)

    # ==================== 学时全自动统计 ====================
    grouped = edited.groupby(["活动分类", "姓名", "学号"]).size().reset_index(name="次数")
    grouped["学时"] = grouped["次数"] * 2

    st.subheader("📊 统计结果")
    st.dataframe(grouped, use_container_width=True)

    # 按（分类 + 学时）分组
    groups = {}
    for _, row in grouped.iterrows():
        key = (row["活动分类"], int(row["学时"]))
        groups.setdefault(key, []).append({"姓名": row["姓名"], "学号": row["学号"]})

    # ==================== 生成规范 Excel ====================
    def make_excel(items):
        wb = Workbook()
        ws = wb.active
        ws["A1"], ws["B1"] = "姓名", "学号"
        yellow = PatternFill(patternType="solid", start_color="FFFF00", end_color="FFFF00")
        center = Alignment(horizontal="center", vertical="center")
        for c in ["A1", "B1"]:
            ws[c].font = Font(name="宋体", size=12)
            ws[c].alignment = center
            ws[c].fill = yellow

        for i, it in enumerate(items, start=2):
            ws.cell(row=i, column=1, value=it["姓名"]).font = Font(name="宋体", size=12)
            ws.cell(row=i, column=2, value=it["学号"]).font = Font(name="宋体", size=12)

        border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=2):
            for cell in row:
                cell.border = border

        ws.column_dimensions["A"].width = 15
        ws.column_dimensions["B"].width = 20

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    st.subheader("📥 下载规范表格")
    for (cat, hrs), items in sorted(groups.items()):
        # 核心修改：文件名加上“第X周”
        week_prefix = f"第{week_num}周" if week_num else ""
        fname = f"{week_prefix}{cat}志愿活动{hrs}学时.xlsx"
        
        # 预览
        st.markdown(f"**预览：{fname}**")
        preview_df = pd.DataFrame(items)
        st.dataframe(preview_df, use_container_width=True, hide_index=True)
        
        st.download_button(
            label=f"⬇️ 下载：{fname}（{len(items)} 人）",
            data=make_excel(items),
            file_name=fname,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key=f"dl_{cat}_{hrs}"
        )
        st.markdown("---")
